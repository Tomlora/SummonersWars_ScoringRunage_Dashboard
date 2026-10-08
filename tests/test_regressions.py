from io import BytesIO
import pytest
from copy import deepcopy
import pandas as pd
from openpyxl import load_workbook
from fonctions.runes import Rune
from fonctions.improvements import grind_opportunities
from fonctions.export import export_excel
from fonctions.leaderboards import visible_players
from fonctions.compare import score_percentile
from fonctions.import_service import analyse_export,validate_export
from pathlib import Path

def test_excel_preserves_input_and_ids():
    frame=pd.DataFrame({'id':[2**40],'value':[14.25]})
    before=frame.copy(deep=True)
    workbook=load_workbook(BytesIO(export_excel(frame,'Index','Runes')))
    assert workbook['Runes'].cell(2,2).value==2**40
    pd.testing.assert_frame_equal(frame,before)

def test_no_stock_means_no_available_gain(export):
    data=grind_opportunities(Rune(export,{}),[])
    assert data['Gain avec stock (max)'].eq(0).all()
    assert data['Meules disponibles'].eq('').all()

def test_stock_respects_set_stat_quality_and_ancient(export):
    rune=deepcopy(export['runes'][0]);export['runes']=[rune]
    code=rune['set_id']*10000+8*100+5
    items=[{'craft_type':2,'craft_type_id':code,'amount':1}]
    normal=grind_opportunities(Rune(export,{}),items)
    assert normal['Gain avec stock (max)'].iloc[0]>0
    rune['rank']=15;rune['extra']=15;rune['class']=16
    ancient=grind_opportunities(Rune(export,{}),items)
    assert ancient['Gain avec stock (max)'].iloc[0]==0

def test_visibility_uses_ids_and_no_guild_is_not_shared():
    data=pd.DataFrame({'id':[1,2,3,4,5],'guilde_id':[0,0,10,10,0],'joueur':['same']*5,'visibility':[0,2,1,4,3]})
    shown=visible_players(data,1,0).set_index('id')
    assert set(shown.index)=={1,3,4,5}
    assert shown.loc[3,'joueur']==shown.loc[4,'joueur']=='***'
    assert shown.loc[1,'joueur']=='same'

def test_percentile_counts_strictly_lower_scores():
    assert score_percentile(pd.DataFrame({'s':[100,100,50]}),'s',100)==pytest.approx(100/3)
    assert score_percentile(pd.DataFrame({'s':[100]}),'s',100)==0

def test_reference_export_scores_and_no_mutation():
    root=Path(__file__).resolve().parents[1]
    data=validate_export(next((root/'SW python').glob('*.json')).read_bytes())
    before=deepcopy(data)
    result=analyse_export(data,pd.DataFrame())
    assert [result[k] for k in ['score','score_spd','score_arte','score_qual']]==[699,513,426,1264]
    assert data==before
    assert len(result['data_rune'].data)==1393
    assert len(result['data_arte'].data_a)==478

def test_detailed_workbook_can_be_read(export):
    from pages_streamlit.optimisation import _build_detailed_workbook, _inventory_dataframe
    runes=Rune(export,{})
    runes.calcul_potentiel();runes.grind()
    book=load_workbook(BytesIO(_build_detailed_workbook(runes,_inventory_dataframe(runes))))
    assert {'Guide','Data_complete','Par rune et monstre','Inventaire'}.issubset(book.sheetnames)
    headers=[cell.value for cell in book['Data_complete'][1]]
    assert len(headers)==len(set(headers))
