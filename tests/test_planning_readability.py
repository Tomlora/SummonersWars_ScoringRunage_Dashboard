from io import BytesIO
import pandas as pd
import pytest
from openpyxl import load_workbook
from fonctions.runes import Rune
from fonctions.planning import grind_plan
from fonctions.grind_efficiency import maximum_grind_efficiencies
from fonctions.workspace import settings
from fonctions.export import export_workbook


def test_slot_filter_allocates_stock_to_selected_runes(export):
    runes=Rune(export,{})
    items=[dict(craft_type=2,craft_type_id=30805,amount=1)]
    original,_=grind_plan(runes,items)
    first=int(original.iloc[0].id_rune)
    slot=int(runes.data.loc[first,'rune_slot'])
    other=3 if slot!=3 else 1
    filtered,remaining=grind_plan(runes,items,slots=[other])
    assert len(filtered)==1 and remaining[2,30805]==0
    assert all(int(runes.data.loc[int(i),'rune_slot'])==other for i in filtered.id_rune)
    assert grind_plan(runes,items,sets=['Will'],slots=[other])[0].empty
    locked=set(runes.data.index[runes.data.rune_slot.astype(int)==other])
    assert grind_plan(runes,items,locked=locked,slots=[other])[0].empty
    assert grind_plan(runes,items,slots=[2])[0].empty  # no SPD substat in this slot in the fixture


@pytest.mark.parametrize('ancient',[False,True])
def test_grind_maxima_keep_existing_better_rolls_and_ungrindable_stats(export,ancient):
    export['runes']=export['runes'][:1]
    rune=export['runes'][0]
    # SPD already has a legendary grind; ATK% can still improve. CRIT and ACC cannot.
    rune['sec_eff']=[[8,25,0,6 if ancient else 5],[4,20,0,0],[9,12,0,0],[12,8,0,0]]
    rune['prefix_eff']=[6,6]
    if ancient:
        rune.update(rank=15,extra=15,**{'class':16})
    runes=Rune(export,{})
    row=runes.data.iloc[0]
    hero,legendary=maximum_grind_efficiencies(runes,row)
    speed=31 if ancient else 30
    expected=lambda attack: round((1+6/40+speed/30+(20+attack)/40+12/30+8/40)/2.8*100,2)
    assert hero==expected(9 if ancient else 7)
    assert legendary==expected(12 if ancient else 10)
    assert legendary>=hero>=row.efficiency


def test_maxima_unavailable_for_incomplete_or_lower_star_runes(export):
    runes=Rune(export,{})
    row=runes.data.iloc[0].copy()
    row['efficiency']=float('nan')
    assert maximum_grind_efficiencies(runes,row)==(None,None)
    row['efficiency']=90
    row['stars']=5
    assert maximum_grind_efficiencies(runes,row)==(None,None)


def test_workbook_keeps_full_readable_data_and_excel_controls():
    plan=pd.DataFrame({'Set':['Swift'],'Slot':[3],'Équipée sur':['=1+1'],'Sous-statistiques':['SPD +30 · ATQ +20%'],'Gain max.':[1.25]})
    original=plan.copy(deep=True)
    data=export_workbook({'Plan':plan,'Stock restant':pd.DataFrame({'Restant':[2]})})
    workbook=load_workbook(BytesIO(data))
    assert workbook.sheetnames==['Plan','Stock restant']
    sheet=workbook['Plan']
    assert [c.value for c in sheet[1]]==list(plan.columns)
    assert sheet['B2'].value==3 and sheet['E2'].value==1.25
    assert sheet['C2'].value=='=1+1' and sheet['C2'].data_type=='s'
    assert sheet.freeze_panes=='A2' and len(sheet.tables)==1
    assert sheet.column_dimensions['A'].width>=14
    pd.testing.assert_frame_equal(plan,original)


@pytest.mark.parametrize('slots',[[0],[7],['2'],[True]])
def test_invalid_slot_preferences_are_rejected(slots):
    with pytest.raises(ValueError):
        settings({'version':1,'filters':{'plan_slots':slots}})
