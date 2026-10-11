from pathlib import Path
import pytest
import pandas as pd
from streamlit.testing.v1 import AppTest
from fonctions.import_service import analyse_export, validate_export, persist_analysis, publish_analysis
from fonctions.workspace import load_settings, settings

ROOT=Path(__file__).resolve().parents[1]


def seed(result,metadata,language='Français'):
    state={};publish_analysis(state,result,metadata)
    at=AppTest.from_file(str(ROOT/'scoring_runage.py'),default_timeout=30)
    for key,value in state.items():at.session_state[key]=value
    at.session_state['translations_selected']=language
    return at.run()


@pytest.mark.parametrize('language',['Français','English'])
def test_plan_search_and_lock_interactions(export,monkeypatch,language):
    monkeypatch.chdir(ROOT);monkeypatch.delenv('API_SQL',raising=False)
    from fonctions import export as excel
    from io import BytesIO
    from openpyxl import load_workbook
    workbooks=[]
    original_export=excel.export_workbook
    def capture(sheets):
        content=original_export(sheets)
        workbooks.append(load_workbook(BytesIO(content)))
        return content
    monkeypatch.setattr(excel,'export_workbook',capture)
    export['rune_craft_item_list']=[dict(craft_type=2,craft_type_id=30805,amount=1)]
    result=analyse_export(validate_export(export),pd.DataFrame())
    at=seed(result,dict(id_joueur=None,report_date='06/10/2026'),language)
    at.switch_page('pages_streamlit/planning.py').run()
    assert not at.exception
    assert 'Principale' not in at.dataframe[0].value.columns
    assert 'Main stat' not in at.dataframe[0].value.columns
    at.toggle(key='ui_plan_details').set_value(True).run()
    shown=at.dataframe[0].value
    owner='Équipée sur' if language=='Français' else 'Equipped on'
    main='Principale' if language=='Français' else 'Main stat'
    subs='Sous-statistiques' if language=='Français' else 'Substats'
    assert {'Set','Slot',owner,main,subs} <= set(shown.columns)
    assert 'Rune' not in shown.columns
    assert shown.iloc[0]['Set']=='Swift'
    assert 'SPD +34' in shown.iloc[0][subs]
    assert any(m.label in ('Gain du plan (points)','Plan gain (points)') for m in at.metric)
    assert any(m.label in ('Max. épique (%)','Max. hero grinds (%)') for m in at.metric)
    assert any(m.label in ('Max. légendaire (%)','Max. legendary grinds (%)') for m in at.metric)
    assert not any(m.value.startswith(('Principale :','Innée :','Main stat :','Innate stat :')) for m in at.markdown)
    at.multiselect(key='plan_slots').set_value([3]).run()
    assert not at.exception
    assert set(at.dataframe[0].value.Slot)=={3}
    sheet=workbooks[-1]['Plan']
    headers=[cell.value for cell in sheet[1]]
    assert subs in headers and main in headers
    assert {row[headers.index('Slot')] for row in sheet.iter_rows(min_row=2,values_only=True)}=={3}
    assert int(result['data_rune'].data.loc[at.selectbox(key='ui_plan_rune').value,'rune_slot'])==3
    at.selectbox(key='plan_sort').set_value('slot').run()
    at.multiselect(key='plan_slots').set_value([]).run()
    assert not at.exception
    search=next(b for b in at.button if b.label in ['Rechercher le build','Search build'])
    search.click().run()
    assert not at.exception
    assert any('SPD +' in s.value for s in at.success)
    assert workbooks[-1].sheetnames==['Build'] and workbooks[-1]['Build'].max_row==7
    at.multiselect(key='locks_runes').set_value(list(result['data_rune'].data.index)).run()
    search=next(b for b in at.button if b.label in ['Rechercher le build','Search build'])
    search.click().run()
    assert not at.exception
    assert any('Aucun build' in s.value or 'No build' in s.value for s in at.info)


def test_preferences_survive_new_session_and_language(engine,export,monkeypatch):
    monkeypatch.chdir(ROOT)
    result=analyse_export(validate_export(export),pd.DataFrame())
    metadata,_=persist_analysis(result)
    at=seed(result,metadata).switch_page('pages_streamlit/optimisation.py').run()
    at.button[0].click().run()
    at.multiselect(key='optimisation_filter_actions').set_value(['gem']).run()
    at.selectbox(key='optimisation_filter_priority').set_value('high').run()
    at.toggle(key='optimisation_show_substats').set_value(True).run()
    at.multiselect(key='optimisation_displayed_substats').set_value(['SPD','ACC']).run()
    at.switch_page('pages_streamlit/planning.py').run()
    at.number_input(key='plan_speed').set_value(40).run()
    at.multiselect(key='plan_slots').set_value([2,4,6]).run()
    at.multiselect(key='locks_runes').set_value([301]).run()
    next(b for b in at.button if b.label=='Enregistrer mes préférences').click().run()
    assert not at.exception
    saved=load_settings(metadata['id_joueur'])
    assert saved['filters']['optimisation_filter_actions']==['gem']
    assert saved['locked_runes']==[301]
    assert saved['filters']['plan_slots']==[2,4,6]
    fresh=seed(result,metadata,'English').switch_page('pages_streamlit/optimisation.py').run()
    fresh.button[0].click().run()
    assert not fresh.exception
    assert fresh.multiselect(key='optimisation_filter_actions').value==['gem']
    assert fresh.selectbox(key='optimisation_filter_priority').value=='high'
    assert fresh.multiselect(key='optimisation_displayed_substats').value==['SPD','ACC']
    fresh.switch_page('pages_streamlit/planning.py').run()
    assert not fresh.exception
    assert fresh.number_input(key='plan_speed').value==40
    assert fresh.multiselect(key='plan_slots').value==[2,4,6]
    assert fresh.multiselect(key='locks_runes').value==[301]


def test_account_switch_clears_local_preferences(export):
    result=analyse_export(validate_export(export),pd.DataFrame())
    state={'compteid':1,'_workspace':settings(),'saved_filter_presets':{'mine':{}},'plan_accuracy':30}
    metadata=dict(id_joueur=None,report_date='06/10/2026')
    publish_analysis(state,result,metadata)
    assert state['plan_accuracy']==30
    result['compteid']=2
    publish_analysis(state,result,metadata)
    assert '_workspace' not in state and 'saved_filter_presets' not in state and 'plan_accuracy' not in state


@pytest.mark.parametrize('language', ['Français', 'English'])
@pytest.mark.parametrize('lookup_case', ['text_index', 'missing'])
def test_planning_sorts_with_text_ids_or_unavailable_details(export, monkeypatch, language, lookup_case):
    monkeypatch.chdir(ROOT)
    monkeypatch.delenv('API_SQL', raising=False)
    export['rune_craft_item_list'] = [dict(craft_type=2, craft_type_id=30805, amount=2)]
    result = analyse_export(validate_export(export), pd.DataFrame())
    if lookup_case == 'text_index':
        result['data_rune'].data.index = result['data_rune'].data.index.astype(str)
    else:
        from fonctions import rune_visual
        identity = rune_visual.rune_identity
        monkeypatch.setattr(rune_visual, 'rune_identity', lambda runes, rune_id: identity(runes, None))
    at = seed(result, dict(id_joueur=None, report_date='06/10/2026'), language)
    at.switch_page('pages_streamlit/planning.py').run()
    for order in ('gain', 'slot', 'set'):
        assert not at.exception
        at.selectbox(key='plan_sort').set_value(order).run()
        assert not at.exception
        shown = at.dataframe[0].value
        assert not shown.empty and {'Set', 'Slot'} <= set(shown)
        if lookup_case == 'missing':
            assert shown.Slot.isna().all() and at.warning
        else:
            assert shown.Slot.notna().all()
            assert at.selectbox(key='ui_plan_rune').options
